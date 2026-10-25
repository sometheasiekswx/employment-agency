"""Content extraction only: never execute macros, scripts or uploaded TeX."""

import io
import re
import zipfile
from html.parser import HTMLParser


class Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip = max(0, self.skip - 1)
        if tag in ("p", "li", "div", "br"):
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def plain(html):
    p = Text()
    p.feed(html)
    return "".join(p.parts).strip()


def extract(name, data):
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("Import limit is 10 MB")
    ext = name.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        from pypdf import PdfReader

        r = PdfReader(io.BytesIO(data))
        if len(r.pages) > 50:
            raise ValueError("Resume exceeds 50 pages")
        text = "\n".join(p.extract_text() or "" for p in r.pages)
    elif ext == "docx":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Expanded DOCX too large")
        from docx import Document

        d = Document(io.BytesIO(data))
        text = "\n".join(
            [p.text for p in d.paragraphs]
            + [c.text for t in d.tables for row in t.rows for c in row.cells]
        )
    elif ext in ("txt", "md", "tex"):
        text = data.decode("utf-8")
        if ext == "tex":
            text = re.sub(r"(?m)(?<!\\)%.*$", "", text)
            text = text.split(r"\begin{document}", 1)[-1]
            text = re.sub(r"\\(?:href)\{[^{}]*\}\{([^{}]*)\}", r"\1", text)
            text = re.sub(r"\\(?:section|item|role|par|newline|newpage)\b", "\n", text)
            text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", " ", text)
            text = text.replace("{", " ").replace("}", " ").replace("\\", " ")
    else:
        raise ValueError("Use PDF, DOCX, TXT, MD or TeX")
    text = text.strip()
    if not text:
        raise ValueError(
            "No text found. For scanned PDFs, supply OCR text or paste your experience."
        )
    if len(text) > 150000:
        raise ValueError("Extracted content too large")
    return text
