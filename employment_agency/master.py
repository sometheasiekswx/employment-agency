"""Optional local master layout. Never compile an imported advert or remote template."""
from pathlib import Path
import os
import re
import shutil
import subprocess
import tempfile
from . import documents, imports
from .store import digest, encode, atomic, now, uid
from .workspaces import inside, settings, city_for

def plain(s):
    s = re.sub(r'\\href\{[^{}]*\}\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\(?:textbf|textit|mbox)\{([^{}]*)\}', r'\1', s)
    s = s.replace(r'\textbar', '|').replace(r'\hfill', '   ')
    s = re.sub(r'\\([$%&#_])', r'\1', s)
    return s.replace('--', '–').strip()


def normalized(s):
    return re.sub(r'\W', '', plain(s)).casefold()


def header_fields(tex):
    """Read identity from the supported master header; fail on layout drift."""
    header = tex.split(r'\begin{document}', 1)[1].split(r'\section{Profile}', 1)[0]
    patterns = {
        'name': r'\\color\{navy\}\s+([^{}]+)\}\\par',
        'headline': r'\{\\large (.*?)\}\\par',
        'location': r'\\mbox\{([^{}]+)\}',
        'email': r'\\href\{mailto:[^{}]+\}\{([^{}]+)\}',
        'phone': r'\\href\{tel:[^{}]+\}\{([^{}]+)\}',
    }
    result = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, header)
        if not match:
            raise ValueError('Cannot read master header field: ' + key)
        result[key] = plain(match.group(1))
    return result


def word_copy(tex, dest):
    """Editable companion with matching content and master-inspired Word styles.
    LaTeX/PDF remain the authoritative layout; Word pagination may differ.
    """
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = Document(); sec=d.sections[0]
    sec.page_height=Mm(297);sec.page_width=Mm(210)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(17)
    d.styles['Normal'].font.name='Latin Modern Roman'; d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.space_after=Pt(2)
    if r'\section{Profile}' not in tex:
        for line in imports.extract('document.tex', tex.encode()).splitlines():
            if line.strip(): d.add_paragraph(line.strip())
        d.save(dest)
        return
    body=tex.split(r'\section{Profile}',1)[1]
    fields = header_fields(tex)
    d.add_heading(fields['name'], 0)
    for line in [fields['headline'], fields['location'], fields['email']+' | '+fields['phone']]:
        d.add_paragraph(line)
    body=r'\section{Profile}'+body
    for line in body.splitlines():
        line=line.strip()
        if not line or line.startswith((r'\begin',r'\end')):continue
        m=re.match(r'\\section\{(.*?)\}',line)
        if m:
            p=d.add_heading(m.group(1),1);p.paragraph_format.keep_with_next=True
            for run in p.runs:run.font.color.rgb=RGBColor.from_string('17384D');run.font.size=Pt(13)
            borders=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom');bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'4');bottom.set(qn('w:color'),'17384D');borders.append(bottom);p._p.get_or_add_pPr().append(borders)
            continue
        m=re.match(r'\\role\{(.*?)\}\{(.*?)\}\{(.*?)\}\{(.*?)\}',line)
        if m:
            role,dates,employer,location=m.groups();p=d.add_paragraph();p.add_run(plain(role)+' | '+plain(dates)).bold=True;p.paragraph_format.keep_with_next=True
            p=d.add_paragraph(employer+' | '+location);p.paragraph_format.keep_with_next=True;continue
        bullet=line.startswith(r'\item ')
        line=line.removeprefix(r'\item ').replace(r'\par','').removesuffix('\\\\')
        d.add_paragraph(plain(line), 'List Bullet' if bullet else None)
    d.save(dest)



def compile_tex(tex, limit=2):
    from pypdf import PdfReader
    compiler = shutil.which('pdflatex')
    if not compiler and Path('/Library/TeX/texbin/pdflatex').is_file():
        compiler = '/Library/TeX/texbin/pdflatex'
    if not compiler:
        raise ValueError('Install a LaTeX compiler to use the master layout')
    # TeX's restricted input/output policy also applies to local templates.
    env = dict(os.environ, openin_any='p', openout_any='p')
    with tempfile.TemporaryDirectory(prefix='career-document-') as temp:
        path = Path(temp)
        (path / 'document.tex').write_text(tex, encoding='utf-8')
        for _ in range(2):
            result = subprocess.run([compiler, '-no-shell-escape', '-interaction=nonstopmode',
                                     '-halt-on-error', 'document.tex'], cwd=path, env=env,
                                    capture_output=True, timeout=60)
            if result.returncode:
                raise ValueError('LaTeX compilation failed; review your local source')
        log = (path / 'document.log').read_text(errors='replace')
        if 'Overfull \\hbox' in log or 'Overfull \\vbox' in log:
            raise ValueError('Document layout overflow; review before using')
        pdf = (path / 'document.pdf').read_bytes()
        reader = PdfReader(path / 'document.pdf')
        if not 1 <= len(reader.pages) <= limit:
            raise ValueError('Document exceeds configured page limit')
        text = '\n'.join(pg.extract_text() or '' for pg in reader.pages)
        return pdf, text


def render(root, directory, job, facts, available_facts=None):
    config = settings(root)
    master = inside(root, config['master'])
    if master.suffix != '.tex':
        raise ValueError('The optional master layout requires a .tex source; use the standard templates for other imports')
    tex = master.read_text()
    city = city_for(job, config)
    tex = tex.replace('[insert city name]', documents.tex_escape(city))
    priorities = {normalized(f['text']): i for i, f in enumerate(facts)}
    def order(match):
        items = re.findall(r'\\item\s+(.*?)(?=\\item|\Z)', match.group(1), re.S)
        items.sort(key=lambda t: priorities.get(normalized(t), 1000))
        return '\\begin{itemize}\n' + ''.join('\\item '+t.strip()+'\n' for t in items) + '\\end{itemize}'
    tex = re.sub(r'\\begin\{itemize\}(.*?)\\end\{itemize\}', order, tex, flags=re.S)
    # Skills come from confirmed workspace facts, never a global personal list.
    skills = []
    requirements = (job.get('description', '') + ' ' + ' '.join(job.get('requirements', []))).casefold()
    for fact in available_facts or facts:
        if fact.get('category') not in ('skill', 'skills') or not fact.get('confirmed') or fact.get('retired') or fact.get('conflict'):
            continue
        for value in re.split(r',\s*|\s+and\s+', fact['text'].split(':', 1)[-1]):
            value = value.strip(' .')
            if value and value.casefold() in requirements and value not in skills:
                skills.append(value)
    tex = tex.replace('[add relevant skills]', documents.tex_escape(', '.join(skills)))
    pdf, text = compile_tex(tex, config['resume_pages'])
    p = Path(directory)
    (p / 'resume.pdf').write_bytes(pdf)
    (p / 'resume.tex').write_text(tex)
    (p / 'resume.txt').write_text(text)
    word_copy(tex, p / 'resume.docx')
    from .packages import rewrite_email
    files = {name: (p / name).read_bytes() for name in documents.FILES}
    rewrite_email(files)
    (p / 'application.eml').write_bytes(files['application.eml'])
    import json
    evidence = json.loads((p / 'evidence.json').read_text())
    evidence['master_template'] = {'source': config['master'], 'sha256': digest(master.read_bytes()),
                                   'layout': 'Local LaTeX master; Word is an editable companion'}
    (p / 'evidence.json').write_bytes(encode(evidence))
    (p / 'checksums.json').write_bytes(encode({n: digest((p / n).read_bytes()) for n in documents.FILES}))


def candidates(path, data):
    """Keep associations explicit for supported role macros; unknown stays blank."""
    if Path(path).suffix.lower() != '.tex':
        return [{'text': line.strip(), 'association': '', 'dates': ''}
                for line in imports.extract(path, data).splitlines() if line.strip()]
    tex = data.decode()
    body = tex.split(r'\begin{document}', 1)[-1]
    # Ignore the decorative header; its contact fields are proposed separately.
    if r'\section{Profile}' in body:
        body = r'\section{Profile}' + body.split(r'\section{Profile}', 1)[1]
    results, association, dates, section = [], '', '', ''
    for line in body.splitlines():
        line = line.strip()
        heading = re.match(r'\\section\*?\{(.*?)\}', line)
        if heading:
            section = heading.group(1).casefold()
            association, dates = '', ''
            continue
        role = re.match(r'\\role\{(.*?)\}\{(.*?)\}\{(.*?)\}\{(.*?)\}', line)
        category = 'skills' if 'skill' in section else 'education' if 'education' in section else 'summary' if section == 'profile' else 'experience'
        if role:
            title, dates, association, location = role.groups()
            results.append({'text': plain(title), 'association': plain(association), 'dates': plain(dates), 'category': category})
            continue
        if not line or line.startswith((r'\begin', r'\end', r'\vspace', '%')):
            continue
        line = line.replace('[add relevant skills]', '').strip()
        try:
            cleaned = imports.extract('line.tex', line.encode()).strip().strip(', ')
        except ValueError:
            continue
        if cleaned:
            results.append({'text': cleaned, 'association': plain(association), 'dates': plain(dates), 'category': category})
    return results or [{'text': line, 'association': '', 'dates': ''} for line in imports.extract(path, data).splitlines() if line.strip()]


def status(workspace):
    config = settings(workspace.path)
    path = inside(workspace.path, config['master'])
    baseline = workspace.read().get('master', {})
    if not path.is_file():
        return {'path': config['master'], 'exists': False, 'changed': bool(baseline.get('sha256'))}
    data = path.read_bytes()
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('Master exceeds 10 MB')
    sha = digest(data)
    return {'path': config['master'], 'exists': True, 'sha256': sha,
            'changed': sha != baseline.get('sha256') or config['master'] != baseline.get('path'), 'accepted_sha256': baseline.get('sha256')}


def propose(workspace):
    info = status(workspace)
    if not info['exists']:
        raise ValueError('Add your master resume to ' + info['path'])
    data = inside(workspace.path, info['path']).read_bytes()
    proposals = candidates(info['path'], data)
    if len(proposals) > 500:
        raise ValueError('Master has more than 500 candidate facts; split it before reconciliation')
    baseline = workspace.read().get('master', {})
    import difflib
    previous = b''
    if baseline.get('sha256'):
        old = inside(workspace.path, '.recovery/masters/' + baseline['sha256'])
        if old.exists():
            previous = old.read_bytes()
    before = imports.extract(info['path'], previous).splitlines() if previous else []
    after = imports.extract(info['path'], data).splitlines()
    identity = {}
    if Path(info['path']).suffix == '.tex':
        try:
            identity = {k: v for k, v in header_fields(data.decode()).items() if k in ('name', 'email', 'phone', 'location')}
        except (ValueError, IndexError):
            pass
    if identity.get('location') and '[insert city name]' in identity['location']:
        identity['location'] = identity['location'].split('|', 1)[0].strip()
    current_state = workspace.read()
    fact_conflicts = []
    for candidate in proposals:
        for existing in current_state['facts']:
            if (existing.get('confirmed') and not existing.get('retired') and
                existing.get('source', {}).get('master') != info['path'] and
                candidate.get('association') and candidate.get('dates') and existing.get('dates') and
                normalized(candidate['association']) == normalized(existing.get('association', '')) and
                normalized(candidate['text']) == normalized(existing['text']) and
                normalized(candidate['dates']) != normalized(existing['dates'])):
                fact_conflicts.append({'fact_id': existing['id'], 'statement': existing['text'],
                                       'existing_dates': existing['dates'], 'proposed_dates': candidate['dates']})
    current_identity = current_state['identity']
    identity = {k: v for k, v in identity.items() if
                (normalized(v) != normalized(current_identity.get(k, '')) if k in ('name', 'phone')
                 else v != current_identity.get(k, ''))}
    return {**info, 'candidates': proposals, 'identity': identity, 'conflicts': fact_conflicts,
            'diff': '\n'.join(difflib.unified_diff(before, after, fromfile='Accepted master', tofile='Current master', lineterm=''))}


def accept(workspace, sha):
    proposal = propose(workspace)
    if sha != proposal['sha256']:
        raise ValueError('Master changed again; refresh and review the new changes')
    if proposal['conflicts']:
        raise ValueError('Resolve the conflicting supplementary fact dates before accepting master updates')
    raw = inside(workspace.path, proposal['path']).read_bytes()
    atomic(inside(workspace.path, '.recovery/masters/' + sha), raw)
    with workspace.edit() as state:
        for fact in state['facts']:
            if fact.get('source', {}).get('master') == proposal['path']:
                fact['retired'], fact['confirmed'] = True, False
        for item in proposal['candidates']:
            state['facts'].append({**item, 'id': uid(), 'category': item.get('category', 'experience'), 'claim_key': '',
                'confirmed': True, 'conflict': False, 'retired': False, 'updated_at': now(),
                'source': {'name': proposal['path'], 'master': proposal['path'], 'sha256': sha}})
        state['identity'].update(proposal.get('identity', {}))
        state['master'] = {'path': proposal['path'], 'sha256': sha, 'accepted_at': now()}
        state['events'].append({'type': 'master_reconciled', 'at': now(), 'sha256': sha})
    return {'confirmed': len(proposal['candidates'])}
